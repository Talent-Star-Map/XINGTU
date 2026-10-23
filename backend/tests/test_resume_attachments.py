"""临时附件和离线身份验证，不读写业务库，不调用大模型。"""
import json
import os
from pathlib import Path
import sys
import tempfile
import subprocess
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
os.environ['DATABASE_URL'] = 'sqlite://'
os.environ['JWT_SECRET'] = 'offline-resume-attachment-secret-at-least-32-characters'
from fastapi import FastAPI
from fastapi.testclient import TestClient
from docx import Document
from PIL import Image
import pypdfium2
from database import create_token
from routers import auth
from services.resume_attachments import word_pdf, remember_filename, metadata

app = FastAPI()
app.include_router(auth.router)
client = TestClient(app)


class AttachmentTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.root = Path(self.folder.name)
        self.files = self.root / 'resumes'
        self.files.mkdir()
        self.patch = patch.object(auth, 'UPLOAD_DIR', str(self.root))
        self.patch.start()
        self.headers = {'Authorization': 'Bearer ' + create_token(11, 'jobseeker')}
        self.path = self.files / '11_import_sample.docx'
        doc = Document()
        doc.add_heading('Attachment preview verification', 0)
        doc.add_paragraph('Python / SQL / 中文简历')
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Project evidence'
        table.cell(0, 1).text = 'Built an API'
        doc.save(self.path)
        self.cache = Path(str(self.path) + '.json')
        self.cache.write_text(json.dumps({'filename': '原始简历.docx', 'skills': ['Python']}), encoding='utf-8')

    def tearDown(self):
        self.patch.stop()
        self.folder.cleanup()

    def get(self, endpoint, filename=None, **kwargs):
        return client.get('/api/auth/' + endpoint,
                          params={'filename': filename or self.path.name, **kwargs}, headers=self.headers)

    def test_list_metadata_and_cached_preview_never_reparse(self):
        with patch('services.resume_parser.parse_resume', side_effect=AssertionError('must not parse')):
            response = self.get('resume-detail')
            listing = client.get('/api/auth/resume-history', headers=self.headers).json()['data']
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['data']['skills'], ['Python'])
        self.assertEqual(listing[0]['filename'], '原始简历.docx')
        self.assertEqual(listing[0]['size'], self.path.stat().st_size)
        self.assertEqual(len(listing), 1)

    def test_original_download_bytes_and_filename_preserved(self):
        response = self.get('resume-file')
        self.assertEqual(response.content, self.path.read_bytes())
        self.assertTrue(response.headers['content-disposition'].startswith('attachment;'))
        self.assertIn('filename*=', response.headers['content-disposition'])
        self.assertEqual(response.headers['cache-control'], 'private, no-store')

    def test_docx_preview_is_real_pdf_with_text_and_table(self):
        response = self.get('resume-file', preview='true')
        self.assertEqual(response.status_code, 200, response.text[:200] if response.status_code != 200 else '')
        self.assertEqual(response.headers['content-type'], 'application/pdf')
        pdf = pypdfium2.PdfDocument(response.content)
        text = pdf[0].get_textpage().get_text_range()
        self.assertIn('Attachment preview verification', text)
        self.assertIn('Project evidence', text)
        self.assertIn('中文简历', text)
        self.assertEqual(self.get('resume-file').content, self.path.read_bytes())

    def test_pdf_and_image_preview_and_text_encoding(self):
        pdf = self.files / '11_example.pdf'
        pdf.write_bytes(word_pdf(self.path))
        response = self.get('resume-file', pdf.name, preview='true')
        self.assertEqual(response.content, pdf.read_bytes())
        self.assertTrue(response.headers['content-disposition'].startswith('inline;'))
        image = self.files / '11_image.png'
        Image.new('RGB', (20, 20), 'white').save(image)
        self.assertEqual(self.get('resume-file', image.name, preview='true').headers['content-type'], 'image/png')
        txt = self.files / '11_note.txt'
        txt.write_bytes('中文技能'.encode('gb18030'))
        self.assertEqual(self.get('resume-file', txt.name, preview='true').text, '中文技能')

    def test_legacy_doc_preview(self):
        from services.doc_convert import _find_soffice
        subprocess.run([_find_soffice(), f'-env:UserInstallation={(self.root / "profile").as_uri()}',
                        '--headless', '--convert-to', 'doc:MS Word 97', '--outdir', str(self.files),
                        str(self.path)], check=True, capture_output=True, timeout=60)
        legacy = self.path.with_suffix('.doc')
        self.assertTrue(legacy.is_file())
        response = self.get('resume-file', legacy.name, preview='true')
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b'%PDF-'))

    def test_paginated_preview_and_page_bounds(self):
        doc = Document(self.path)
        doc.add_page_break()
        doc.add_paragraph('Second page')
        doc.save(self.path)
        response = self.get('resume-file', preview='true', page=2)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers['content-type'], 'image/png')
        self.assertEqual(response.headers['x-page-count'], '2')
        self.assertTrue(response.content.startswith(b'\x89PNG'))
        self.assertEqual(self.get('resume-file', preview='true', page=3).status_code, 400)
        self.assertEqual(self.get('resume-file', preview='true', page=0).status_code, 422)
        tiff = self.files / '11_image.tiff'
        Image.new('RGB', (20, 20), 'white').save(tiff, save_all=True, append_images=[Image.new('RGB', (20, 20), 'blue')])
        response = self.get('resume-file', tiff.name, preview='true', page=2)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers['x-page-count'], '2')

    def test_upload_returns_attachment_and_list_includes_it(self):
        with patch('services.resume_parser.parse_resume', return_value={'success': True, 'data': {'skills': ['Python']}}):
            response = client.post('/api/auth/resume-parse', headers=self.headers,
                                   files={'file': ('上传简历.docx', self.path.read_bytes(), 'application/octet-stream')})
        self.assertEqual(response.status_code, 200)
        attachment = response.json()['data']['attachment']
        self.assertEqual(attachment['filename'], '上传简历.docx')
        self.assertTrue(attachment['path'].startswith('11_'))
        self.assertTrue(Path(str(self.files / attachment['path']) + '.metadata').is_file())
        listing = client.get('/api/auth/resume-history', headers=self.headers).json()['data']
        self.assertEqual(len(listing), 2)

    def test_original_name_survives_missing_or_overwritten_parse_cache(self):
        remember_filename(self.path, '本地原始名称（最终版）.docx')
        self.cache.write_text(json.dumps({'filename': self.path.name}), encoding='utf-8')
        self.assertEqual(metadata(self.path)['filename'], '本地原始名称（最终版）.docx')
        self.cache.unlink()
        self.assertEqual(metadata(self.path)['filename'], '本地原始名称（最终版）.docx')
        response = self.get('resume-file')
        self.assertIn('filename*=', response.headers['content-disposition'])

    def test_owner_role_path_and_cache_access_are_enforced(self):
        for name in ['12_other.docx', '../11_import_sample.docx', r'folder\11_import_sample.docx']:
            self.assertEqual(self.get('resume-file', name).status_code, 403)
            self.assertEqual(self.get('resume-detail', name).status_code, 403)
        self.assertEqual(self.get('resume-file', self.path.name + '.json').status_code, 400)
        self.assertEqual(self.get('resume-file', '11_missing.pdf').status_code, 404)
        self.assertEqual(client.get('/api/auth/resume-detail', params={'filename': self.path.name}).status_code, 401)
        self.headers = {'Authorization': 'Bearer ' + create_token(11, 'enterprise')}
        self.assertEqual(self.get('resume-file').status_code, 403)

    def test_no_cache_still_allows_attachment_viewing(self):
        self.cache.unlink()
        data = self.get('resume-detail').json()['data']
        self.assertEqual(data['attachment']['filename'], self.path.name)
        self.assertFalse(data['attachment']['parsed_available'])

    def test_reparse_preserves_original_filename_and_delete_removes_both(self):
        with patch('services.resume_parser.parse_resume', return_value={'success': True, 'data': {'skills': ['SQL']}}):
            response = client.post('/api/auth/resume-reparse', params={'filename': self.path.name}, headers=self.headers)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(json.loads(self.cache.read_text())['filename'], '原始简历.docx')
        response = client.delete('/api/auth/resume-delete', params={'filename': self.path.name}, headers=self.headers)
        self.assertTrue(response.json()['success'])
        self.assertFalse(self.path.exists())
        self.assertFalse(self.cache.exists())
        self.assertFalse(Path(str(self.path) + '.metadata').exists())


if __name__ == '__main__':
    unittest.main()
