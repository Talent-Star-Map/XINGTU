"""浏览器验收：所有 API 使用本机测试附件，禁止访问真实业务数据。"""
import json
import os
import re
from pathlib import Path
import tempfile
import sys
from urllib.parse import urlparse, parse_qs

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from playwright.sync_api import sync_playwright, expect
from services.resume_attachments import word_pdf, pdf_page


def main():
    with tempfile.TemporaryDirectory() as folder:
        path = Path(folder) / '11_preview.docx'
        doc = Document()
        doc.add_heading('Resume attachment preview', 0)
        doc.add_paragraph('Python / SQL / 中文简历')
        doc.add_page_break()
        doc.add_paragraph('Second page: project experience')
        doc.save(path)
        pdf = word_pdf(path)
        info = {'path': path.name, 'filename': '测试简历.docx', 'extension': 'docx', 'size': path.stat().st_size,
                'date': '2026-10-23 12:00', 'skills': ['Python', 'SQL'], 'skill_count': 2}
        parsed = {'name': '附件预览测试', 'school': '测试学校', 'skills': ['Python', 'SQL'],
                  'target_position': 'Python开发工程师', 'bio': '本地附件预览测试', 'attachment': info}

        def mock(route):
            request = route.request
            endpoint = urlparse(request.url).path
            if not endpoint.startswith('/api/'):
                upstream = route.fetch(headers={**request.headers, 'host': 'localhost'})
                route.fulfill(response=upstream)
                return
            if endpoint == '/api/auth/resume-file':
                inline = 'preview=true' in request.url
                if inline:
                    number = int(parse_qs(urlparse(request.url).query).get('page', ['1'])[0])
                    image, count = pdf_page(pdf, number)
                    route.fulfill(body=image, content_type='image/png', headers={'X-Page-Count': str(count)})
                    return
                route.fulfill(body=pdf if inline else path.read_bytes(), content_type='application/pdf' if inline else 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
                return
            data = []
            if endpoint == '/api/auth/resume-history': data = [info]
            elif endpoint == '/api/auth/resume-detail':
                if parse_qs(urlparse(request.url).query).get('filename') == ['11_missing.pdf']:
                    route.fulfill(status=404, json={'detail': '原附件不存在'})
                    return
                data = parsed
            elif endpoint == '/api/auth/profile': data = {'username': '附件测试', 'skills': 'Python,SQL'}
            route.fulfill(json={'success': True, 'data': data})

        with sync_playwright() as p:
            browser = p.chromium.launch(executable_path=os.getenv('CHROMIUM_PATH', '/usr/bin/chromium'), headless=True, args=['--no-sandbox'])
            page = browser.new_page(viewport={'width': 1440, 'height': 1100}, accept_downloads=True)
            page.route('http://frontend:5173/**', mock)
            page.add_init_script("localStorage.setItem('xingtu_token','offline-ui-test');localStorage.setItem('xingtu_role','jobseeker');localStorage.setItem('xingtu_user',JSON.stringify({id:11,username:'附件测试',role:'jobseeker'}));")
            page.goto('http://frontend:5173/login/jobseeker#/resume')
            expect(page.get_by_text('简历附件 (1)', exact=True)).to_be_visible()
            with page.expect_download() as download:
                page.get_by_title('下载原附件', exact=True).click()
            assert download.value.suggested_filename == info['filename']
            assert Path(download.value.path()).read_bytes() == path.read_bytes()
            page.get_by_title('预览', exact=True).click()
            expect(page.get_by_text('简历预览', exact=True)).to_be_visible()
            dialog = page.get_by_role('dialog', name='简历预览')
            assert dialog.bounding_box()['width'] <= 577, 'preview dialog must retain original width'
            expect(page.get_by_text('原文预览', exact=True)).to_have_count(0)
            page.screenshot(path='/tmp/resume-attachments-ui.png', full_page=True)
            page.evaluate('window.__attachmentEntryMarker = true')
            page.get_by_role('link', name='预览附件', exact=True).click()
            expect(page).to_have_url(re.compile(r'/resume-reader\?filename='))
            expect(page.get_by_role('dialog')).to_have_count(0)
            expect(page.get_by_text('工作台', exact=True)).to_have_count(0)
            frame = page.get_by_role('img', name='测试简历.docx 第1页', exact=True)
            expect(frame).to_be_visible(timeout=15000)
            assert page.evaluate('typeof window.__attachmentEntryMarker') == 'undefined', 'reader must load a fresh document'
            expect(frame).to_have_js_property('complete', True)
            assert frame.evaluate('(image) => image.naturalWidth') > 0
            expect(page.get_by_role('spinbutton', name='页码')).to_have_value('1')
            page.screenshot(path='/tmp/resume-attachment-preview-ui.png', full_page=True)
            page.get_by_role('button', name='下一页', exact=True).click()
            expect(page.get_by_role('img', name='测试简历.docx 第2页', exact=True)).to_be_visible()
            expect(page).to_have_url(re.compile(r'page=2'))
            page.get_by_role('combobox', name='缩放比例').select_option('100')
            original_width = page.get_by_test_id('resume-paper').bounding_box()['width']
            page.get_by_role('button', name='放大', exact=True).click()
            assert page.get_by_test_id('resume-paper').bounding_box()['width'] > original_width
            with page.expect_download() as download:
                page.get_by_role('button', name='下载原文件', exact=True).click()
            assert download.value.suggested_filename == info['filename']
            page.reload()
            expect(page.get_by_role('img', name='测试简历.docx 第2页', exact=True)).to_be_visible()
            page.set_viewport_size({'width': 390, 'height': 844})
            expect(page.get_by_role('complementary', name='页码导航')).not_to_be_visible()
            assert page.get_by_test_id('resume-paper').bounding_box()['width'] <= 390
            page.get_by_role('button', name='上一页', exact=True).click()
            expect(page.get_by_role('img', name='测试简历.docx 第1页', exact=True)).to_be_visible()
            page.screenshot(path='/tmp/resume-reader-mobile-ui.png', full_page=True)
            page.set_viewport_size({'width': 1440, 'height': 1100})
            page.get_by_role('link', name='返回简历管理', exact=True).click()
            expect(page.get_by_text('简历附件 (1)', exact=True)).to_be_visible()
            page.goto('http://frontend:5173/resume-reader?filename=11_missing.pdf')
            expect(page.get_by_role('alert')).to_contain_text('原附件不存在')
            page.unroute_all(behavior='wait')
            browser.close()
            print(json.dumps({'list_and_download': 'passed', 'original_dialog_width': 'passed', 'reader_route': 'passed',
                              'page_and_zoom': 'passed', 'reload_and_return': 'passed', 'mobile': 'passed', 'missing_attachment': 'passed'}))


if __name__ == '__main__':
    main()
